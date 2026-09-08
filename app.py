from flask import Flask, render_template, request, jsonify, send_from_directory, send_file
from werkzeug.exceptions import HTTPException
import os
import shutil
import itertools
import logging
import time
from io import BytesIO
from datetime import datetime

from PIL import Image
import torch
from torchvision import transforms
from pytorch_msssim import ms_ssim
from ultralytics import YOLO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from custom_zipfile import ZipFile

logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s')
logger = logging.getLogger(__name__)

app = Flask(__name__, static_folder='static')

# 配置上传文件存储路径
UPLOAD_FOLDER = 'uploads'
EXTRACT_FOLDER = 'extracted'
ALLOWED_EXTENSIONS = {'zip'}
ZIP_FILENAME = 'upload.zip'

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 500 * 1024 * 1024  # 500MB

# 确保上传和解压目录存在
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(EXTRACT_FOLDER, exist_ok=True)

# 全局分析结果（上传后填充），初始化为空列表避免未上传时访问 /export_excel 等路由报 NameError
results = []

# 分类结果中文标签
CLASSIFICATION_LABELS = {
    'perfect': '无异常',
    'slightly-stained': '略有脏污',
    'awfully-stained': '严重脏污',
    'offline': '离线',
}

# 加载模型
MODEL_PATH = 'models/best.pt'
model = YOLO(MODEL_PATH)

# 性能参数
PREDICT_BATCH_SIZE = 32   # YOLO 批量推理的批大小
FREEZE_THRESHOLD = 0.997  # MS-SSIM 判定卡住阈值
MS_SSIM_PAIR_BATCH = 256  # MS-SSIM 批量计算的图片对数量

# 判断是否是允许的文件类型
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def _rmtree_with_retry(path, retries=3):
    """Windows 下文件句柄释放有延迟，删除失败时短暂等待后重试"""
    for attempt in range(retries):
        try:
            shutil.rmtree(path)
            return
        except PermissionError:
            if attempt == retries - 1:
                raise
            time.sleep(0.2)

def clean_folders():
    """清空上传与解压目录"""
    for folder in [UPLOAD_FOLDER, EXTRACT_FOLDER]:
        if os.path.exists(folder):
            _rmtree_with_retry(folder)
        os.makedirs(folder)

def classify_single(image_path):
    try:
        result = model.predict(image_path, task='classify', verbose=False)[0]
        return result.names[result.probs.top1]
    except Exception:
        logger.exception('图片分类失败: %s', image_path)
        return 'unknown'

def classify_images(image_paths):
    """批量分类图片；某批失败时回退为逐张处理，避免单张坏图拖垮整批"""
    classifications = []
    for start in range(0, len(image_paths), PREDICT_BATCH_SIZE):
        batch = image_paths[start:start + PREDICT_BATCH_SIZE]
        try:
            batch_results = model.predict(batch, task='classify', verbose=False)
            classifications.extend(r.names[r.probs.top1] for r in batch_results)
        except Exception:
            logger.exception('批量推理失败，回退为逐张处理')
            classifications.extend(classify_single(p) for p in batch)
    return classifications

# 图像相似度检测相关
transform = transforms.Compose([
    transforms.Resize((256, 256)),
    transforms.ToTensor(),
])

def check_camera_freeze(image_paths):
    """检查摄像头是否卡住：所有图片两两 MS-SSIM 均不低于阈值视为卡住"""
    if len(image_paths) <= 1:
        return False

    # 每张图片只加载一次，后续成对比较复用张量
    tensors = []
    for path in image_paths:
        try:
            img = Image.open(path).convert('RGB')
            tensors.append(transform(img))
        except Exception:
            logger.exception('读取图片失败: %s', path)
            return False

    pair_indices = list(itertools.combinations(range(len(tensors)), 2))
    for start in range(0, len(pair_indices), MS_SSIM_PAIR_BATCH):
        chunk = pair_indices[start:start + MS_SSIM_PAIR_BATCH]
        a = torch.stack([tensors[i] for i, _ in chunk]) * 255
        b = torch.stack([tensors[j] for _, j in chunk]) * 255
        try:
            vals = ms_ssim(a, b, data_range=255, size_average=False)
        except Exception:
            logger.exception('MS-SSIM 计算失败')
            return False
        if (vals < FREEZE_THRESHOLD).any():
            return False
    return True

def calculate_health_grade(classifications, is_frozen):
    """计算健康度等级"""
    if is_frozen:
        return 'D', True  # 返回等级和是否因为卡住而判D

    if all(c == 'perfect' for c in classifications):
        return 'A', False
    elif 'offline' in classifications or all(c == 'awfully-stained' for c in classifications):
        return 'D', False
    elif 'awfully-stained' in classifications:
        return 'C', False
    elif 'slightly-stained' in classifications:
        return 'B', False
    return 'A', False

def build_date_statistics(all_results):
    """汇总所有日期的健康度分布，返回 (排序后的日期列表, 统计字典)"""
    date_statistics = {}
    for result in all_results:
        for date_result in result['dates']:
            stats = date_statistics.setdefault(
                date_result['date'], {'A': 0, 'B': 0, 'C': 0, 'D': 0, 'total': 0})
            stats[date_result['health_grade']] += 1
            stats['total'] += 1
    all_dates = sorted(date_statistics.keys())
    return all_dates, date_statistics

@app.errorhandler(413)
def request_too_large(e):
    return jsonify({'error': '文件大小超过限制（500MB）'}), 413

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    try:
        if 'file' not in request.files:
            return jsonify({'error': '没有文件上传'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': '没有选择文件'}), 400

        if not allowed_file(file.filename):
            return jsonify({'error': '不支持的文件类型'}), 400

        try:
            # 清理旧文件，避免上一次上传的数据残留混入本次结果
            clean_folders()

            # 固定文件名保存，避免上传文件名中的特殊字符影响存储路径
            zip_path = os.path.join(app.config['UPLOAD_FOLDER'], ZIP_FILENAME)
            file.save(zip_path)

            # 解压文件
            with ZipFile(zip_path, 'r') as zip_ref:
                zip_ref.extractall(EXTRACT_FOLDER)

            # 解压完成即可删除压缩包，释放磁盘空间
            os.remove(zip_path)

            # 遍历目录结构，收集待分类图片
            temp_results = {}
            all_image_paths = []
            for date_folder in os.listdir(EXTRACT_FOLDER):
                date_path = os.path.join(EXTRACT_FOLDER, date_folder)
                if not os.path.isdir(date_path):
                    continue
                for camera_folder in os.listdir(date_path):
                    camera_path = os.path.join(date_path, camera_folder)
                    if not os.path.isdir(camera_path):
                        continue
                    entry = temp_results.setdefault(camera_folder, {}).setdefault(date_folder, {
                        'classifications': [],
                        'images': [],
                        'image_paths': []
                    })
                    for image_file in os.listdir(camera_path):
                        if image_file.lower().endswith(('.png', '.jpg', '.jpeg')):
                            image_path = os.path.join(camera_path, image_file)
                            all_image_paths.append(image_path)
                            entry['images'].append({'image': image_file, 'classification': None})
                            entry['image_paths'].append(image_path)

            # 批量分类所有图片后按序回填
            all_classifications = classify_images(all_image_paths)
            idx = 0
            for dates in temp_results.values():
                for data in dates.values():
                    for image in data['images']:
                        image['classification'] = all_classifications[idx]
                        idx += 1
                    data['classifications'] = [img['classification'] for img in data['images']]

            # 计算每个摄像头在每个日期的健康度等级
            global results
            results = []
            for camera, dates in temp_results.items():
                camera_results = {
                    'camera': camera,
                    'dates': []
                }

                for date, data in dates.items():
                    is_frozen = check_camera_freeze(data['image_paths'])
                    health_grade, frozen_cause = calculate_health_grade(data['classifications'], is_frozen)
                    camera_results['dates'].append({
                        'date': date,
                        'health_grade': health_grade,
                        'images': data['images'],
                        'frozen': frozen_cause
                    })

                # 按日期排序
                camera_results['dates'].sort(key=lambda x: x['date'])
                results.append(camera_results)

            # 按摄像头名称排序
            results.sort(key=lambda x: x['camera'])

            all_dates, date_statistics = build_date_statistics(results)

            return render_template('results.html',
                                   results=results,
                                   dates=all_dates,
                                   date_statistics=date_statistics,
                                   labels=CLASSIFICATION_LABELS)

        except HTTPException:
            # 让 413 等 HTTP 错误走 Flask 全局错误处理器，返回统一 JSON
            raise
        except Exception as e:
            clean_folders()
            return jsonify({'error': f'处理文件时出错: {str(e)}'}), 500

    except HTTPException:
        raise
    except Exception as e:
        clean_folders()
        return jsonify({'error': f'服务器错误: {str(e)}'}), 500

# 添加新的路由来访问图片
@app.route('/image/<path:image_path>')
def serve_image(image_path):
    # 从路径中获取目录名和文件名
    directory = os.path.dirname(image_path)
    filename = os.path.basename(image_path)
    return send_from_directory(os.path.join(EXTRACT_FOLDER, directory), filename)

# 添加新的路由来处理清理
@app.route('/cleanup', methods=['POST'])
def cleanup():
    try:
        clean_folders()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# 添加新的路由处理函数
@app.route('/review', methods=['POST'])
def review():
    try:
        camera = request.form['camera']
        date = request.form['date']
        frozen = 'frozen' in request.form
        classifications = {}

        # 处理图片分类结果
        for key, value in request.form.items():
            if key.startswith('classifications['):
                image_name = key[len('classifications['):-1]  # 提取图片名称
                classifications[image_name] = value

        # 更新结果
        for result in results:
            if result['camera'] != camera:
                continue
            for date_result in result['dates']:
                if date_result['date'] != date:
                    continue
                date_result['frozen'] = frozen
                for image in date_result['images']:
                    if image['image'] in classifications:
                        image['classification'] = classifications[image['image']]

                # 重新计算健康度等级
                new_classifications = [img['classification'] for img in date_result['images']]
                health_grade, _ = calculate_health_grade(new_classifications, frozen)
                date_result['health_grade'] = health_grade
            break

        all_dates, date_statistics = build_date_statistics(results)

        return render_template('results.html',
                               results=results,
                               dates=all_dates,
                               date_statistics=date_statistics,
                               labels=CLASSIFICATION_LABELS)

    except Exception as e:
        return jsonify({'error': f'处理审核结果时出错: {str(e)}'}), 500

# 添加新的路由函数
@app.route('/export_excel')
def export_excel():
    try:
        if not results:
            return jsonify({'error': '暂无分析结果，请先上传数据'}), 400

        wb = Workbook()
        ws = wb.active
        ws.title = "摄像头健康度分析"

        # 设置表头
        headers = ['摄像头', '日期', '健康度等级', '是否卡住', '图片名称', '图片状态']
        ws.append(headers)

        # 预处理数据，按摄像头分组并排序
        camera_groups = {}
        for result in results:
            camera = result['camera']
            if camera not in camera_groups:
                camera_groups[camera] = []
            for date_result in result['dates']:
                camera_groups[camera].append({
                    'date': date_result['date'],
                    'health_grade': date_result['health_grade'],
                    'frozen': date_result['frozen'],
                    'images': date_result['images']
                })

        # 按摄像头名称排序
        sorted_cameras = sorted(camera_groups.keys())

        current_row = 2  # 从第2行开始写入数据
        for camera in sorted_cameras:
            camera_start_row = current_row
            dates = camera_groups[camera]

            # 按日期排序
            dates.sort(key=lambda x: x['date'])

            for date_data in dates:
                # 获取该组数据的图片数量
                image_count = len(date_data['images'])

                # 写入第一行数据
                first_row = [
                    camera if date_data == dates[0] else '',  # 只在摄像头的第一行写入摄像头名称
                    date_data['date'],
                    date_data['health_grade'],
                    '是' if date_data['frozen'] else '否'
                ]

                # 处理每张图片
                for idx, image in enumerate(date_data['images']):
                    status = CLASSIFICATION_LABELS.get(image['classification'], image['classification'])
                    if idx == 0:
                        # 第一张图片与前面的数据在同一行
                        ws.append(first_row + [image['image'], status])
                    else:
                        # 其他图片只写入图片相关信息
                        ws.append(['', '', '', '', image['image'], status])

                # 如果有多张图片，合并当前日期的前四列单元格
                if image_count > 1:
                    for col in ['B', 'C', 'D']:  # 不包括A列，因为A列要跨所有日期合并
                        ws.merge_cells(f'{col}{current_row}:{col}{current_row + image_count - 1}')

                current_row += image_count

            # 合并整个摄像头的A列
            camera_end_row = current_row - 1
            if camera_end_row > camera_start_row:
                ws.merge_cells(f'A{camera_start_row}:A{camera_end_row}')

        # 设置样式
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = header_font

        # 设置所有单元格的对齐方式
        for row in ws.rows:
            for cell in row:
                cell.alignment = Alignment(
                    vertical='center',
                    horizontal='left',
                    wrapText=True
                )

        # 预设最小列宽
        min_widths = {
            'A': 25,  # 摄像头
            'B': 12,  # 日期
            'C': 12,  # 健康度等级
            'D': 10,  # 是否卡住
            'E': 20,  # 图片名称
            'F': 15   # 图片状态
        }

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            column = list(col)
            column_letter = column[0].column_letter

            # 计算该列中最长的内容
            for cell in column:
                try:
                    if cell.value:
                        # 计算实际显示长度（中文字符计为2个单位）
                        length = sum(2 if ord(char) > 127 else 1 for char in str(cell.value))
                        max_length = max(max_length, length)
                except Exception:
                    pass

            # 确保不小于预设最小宽度
            min_width = min_widths.get(column_letter, 10)
            adjusted_width = max(max_length + 2, min_width)

            ws.column_dimensions[column_letter].width = adjusted_width

        # 设置冻结窗格（固定表头）
        ws.freeze_panes = 'A2'

        # 保存到内存中
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        # 生成文件名
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'摄像头健康评估结果_{timestamp}.xlsx'

        return send_file(
            excel_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        return jsonify({'error': f'导出Excel时出错: {str(e)}'}), 500

if __name__ == '__main__':
    app.run(port=5000, debug=False)
